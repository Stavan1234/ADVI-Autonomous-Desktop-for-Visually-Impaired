from docx import Document

doc = Document()
doc.add_heading('Top Secret Mission Briefing', 0)
doc.add_paragraph('Operation Midnight Falcon is a go. The primary target is located at the coordinates 45.9, -12.4.')
doc.save('test_document.docx')
print("Created test_document.docx")
